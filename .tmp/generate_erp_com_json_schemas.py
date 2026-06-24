import json
import os
import openpyxl
from collections import OrderedDict

xlsx_path = "c:/Users/HP/Documents/repo_git/terraform-gcp/Diccionario (2).xlsx"
out_dir = "c:/Users/HP/Documents/repo_git/terraform-gcp/environments/dev/schemas/bronze/erp_com"
os.makedirs(out_dir, exist_ok=True)

ws = openpyxl.load_workbook(xlsx_path, data_only=True)["ERP Com"]


def bq_type(src_type, width, scale):
    t = (src_type or "").strip().lower()
    try:
        w = int(width) if width is not None and str(width).strip() != "" else None
    except Exception:
        w = None
    try:
        s = int(scale) if scale is not None and str(scale).strip() != "" else 0
    except Exception:
        s = 0

    if t in ("varchar", "char", "text", "string"):
        return "STRING"
    if t == "date":
        return "DATE"
    if t in ("datetime", "timestamp"):
        return "TIMESTAMP"
    if t in ("int", "integer", "smallint", "bigint"):
        return "INT64"
    if t in ("numeric", "decimal", "number"):
        if s and s > 0:
            return "NUMERIC"
        if w is not None and w <= 18:
            return "INT64"
        return "NUMERIC"
    if t in ("float", "double", "real"):
        return "FLOAT64"
    if t in ("boolean", "bool"):
        return "BOOL"
    return "STRING"


schemas = OrderedDict()
for r in range(2, ws.max_row + 1):
    table = ws.cell(r, 4).value
    col = ws.cell(r, 5).value
    src_t = ws.cell(r, 6).value
    width = ws.cell(r, 7).value
    scale = ws.cell(r, 8).value
    nulls = ws.cell(r, 9).value

    if not table or not col:
        continue

    table = str(table).strip().lower()
    col = str(col).strip()
    mode = "NULLABLE" if str(nulls).strip().upper() == "Y" else "REQUIRED"

    schemas.setdefault(table, [])
    schemas[table].append(
        {
            "name": col,
            "type": bq_type(src_t, width, scale),
            "mode": mode,
        }
    )

for table, fields in schemas.items():
    path = os.path.join(out_dir, f"{table}.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(fields, f, ensure_ascii=True, indent=2)

print(f"Generated {len(schemas)} schema files in {out_dir}")
print("Tables:", ", ".join(schemas.keys()))
