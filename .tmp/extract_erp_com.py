import openpyxl
from collections import OrderedDict

path = "c:/Users/HP/Documents/repo_git/terraform-gcp/Diccionario (2).xlsx"
ws = openpyxl.load_workbook(path, data_only=True)["ERP Com"]


def bq_type(src_type, width, scale):
    t = (src_type or "").strip().lower()
    try:
        w = int(width) if width is not None and str(width).strip() != "" else None
    except Exception:
        w = None
    try:
        s = int(scale) if scale is not None and str(scale).strip() != "" else 0
    except Exception:
        s = 0

    if t in ("varchar", "char", "text", "string"):
        return "STRING"
    if t == "date":
        return "DATE"
    if t in ("datetime", "timestamp"):
        return "TIMESTAMP"
    if t in ("int", "integer", "smallint", "bigint"):
        return "INT64"
    if t in ("numeric", "decimal", "number"):
        if s and s > 0:
            return "NUMERIC"
        if w is not None and w <= 18:
            return "INT64"
        return "NUMERIC"
    if t in ("float", "double", "real"):
        return "FLOAT64"
    if t in ("boolean", "bool"):
        return "BOOL"
    return "STRING"


schemas = OrderedDict()
for r in range(2, ws.max_row + 1):
    table = ws.cell(r, 4).value
    col = ws.cell(r, 5).value
    src_t = ws.cell(r, 6).value
    width = ws.cell(r, 7).value
    scale = ws.cell(r, 8).value
    nulls = ws.cell(r, 9).value

    if not table or not col:
        continue

    table = str(table).strip().lower()
    col = str(col).strip()
    mode = "NULLABLE" if str(nulls).strip().upper() == "Y" else "REQUIRED"

    schemas.setdefault(table, [])
    schemas[table].append(
        {
            "name": col,
            "type": bq_type(src_t, width, scale),
            "mode": mode,
        }
    )

print(f"TABLE_COUNT={len(schemas)}")
print("TABLE_NAMES=" + ",".join(schemas.keys()))

lines = []
lines.append("external_tables = {")
for table, fields in schemas.items():
    lines.append(f"  {table} = {{")
    lines.append('    dataset_id  = "gim_dataset_brz_dev"')
    lines.append('    source      = "erp_com"')
    lines.append(
        f"    source_uris = [\"gs://gim-lakehouse-bronze-${{var.environment}}/erp_com/{table}/year=*/month=*/day=*/*.parquet\"]"
    )
    lines.append("    schema = [")
    for f in fields:
        lines.append("      {")
        lines.append(f"        name = \"{f['name']}\"")
        lines.append(f"        type = \"{f['type']}\"")
        lines.append(f"        mode = \"{f['mode']}\"")
        lines.append("      },")
    lines.append("    ]")
    lines.append("  }")
lines.append("}")

output_path = "c:/Users/HP/Documents/repo_git/terraform-gcp/.tmp/erp_com_external_tables_dev.auto.tfvars"
with open(output_path, "w", encoding="utf-8") as f:
    f.write("\n".join(lines) + "\n")

print("WROTE=" + output_path)
